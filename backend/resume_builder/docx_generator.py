from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .resume_data import (
    csv_list,
    entry_has_content,
    format_dates,
    normalize_resume,
    skill_list,
)
from .template_engine import get_template


def _set_run_color(run, hex_color):
    hex_value = (hex_color or "111827").replace("#", "")
    if len(hex_value) != 6:
        hex_value = "111827"
    run.font.color.rgb = RGBColor(
        int(hex_value[0:2], 16),
        int(hex_value[2:4], 16),
        int(hex_value[4:6], 16),
    )


def _add_bottom_border(paragraph, color_hex):
    p_el = paragraph._p
    p_pr = p_el.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color_hex.replace("#", ""))
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def _add_hyperlink(paragraph, text, url):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1d4ed8")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    new_run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    new_run.append(text_el)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def generate_docx(resume, output_file):
    if isinstance(resume, dict) and "sections" in resume:
        return _generate_legacy_docx(resume, output_file)

    data = normalize_resume(resume)
    theme = get_template(data["template"])
    doc = Document()

    wide = bool(theme.get("wide"))
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.left_margin = Inches(0.75 if wide else 0.65)
        section.right_margin = Inches(0.75 if wide else 0.65)
        section.top_margin = Inches(0.6 if wide else 0.55)
        section.bottom_margin = Inches(0.55)

    header_color = theme["header"]
    if header_color in ("#f8fafc",):
        header_color = "#0f172a"

    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = name.add_run(data.get("full_name") or "Resume")
    run.bold = True
    run.font.size = Pt(22)
    run.font.name = "Times New Roman" if theme.get("serif") else "Calibri"
    _set_run_color(run, header_color)

    if data.get("professional_title"):
        role = doc.add_paragraph()
        role.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = role.add_run(data["professional_title"])
        run.font.size = Pt(12)
        _set_run_color(run, theme["accent"])

    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    first = True
    for key, label in (
        ("email", None),
        ("phone", None),
        ("location", None),
        ("linkedin", "LinkedIn"),
        ("github", "GitHub"),
        ("portfolio", "Portfolio"),
    ):
        value = data.get(key)
        if not value:
            continue
        if not first:
            sep = contact.add_run("  |  ")
            sep.font.size = Pt(9)
        first = False
        if key in ("linkedin", "github", "portfolio") and str(value).startswith("http"):
            _add_hyperlink(contact, label, value)
        else:
            run = contact.add_run(value)
            run.font.size = Pt(9)
            muted = theme["muted"] if theme["muted"] != "#cbd5e1" else "#475569"
            _set_run_color(run, muted)
    _add_bottom_border(contact, theme["rule"])

    if data.get("summary"):
        _heading(doc, "Professional Summary", theme)
        doc.add_paragraph(data["summary"])

    order = theme.get("order")
    sequence = (
        ("education", "projects", "experience")
        if order == "graduate"
        else ("experience", "projects", "education")
    )

    for key in sequence:
        if key == "experience":
            blocks = [item for item in data.get("experience") or [] if entry_has_content(item)]
            if not blocks:
                continue
            _heading(doc, "Work Experience", theme)
            for item in blocks:
                title = " • ".join(
                    part for part in (item.get("role"), item.get("company")) if part
                )
                paragraph = doc.add_paragraph()
                run = paragraph.add_run(title)
                run.bold = True
                meta = "  |  ".join(
                    bit
                    for bit in (
                        format_dates(
                            item.get("start"), item.get("end"), item.get("current")
                        ),
                        item.get("location") or "",
                    )
                    if bit
                )
                if meta:
                    meta_p = doc.add_paragraph()
                    meta_run = meta_p.add_run(meta)
                    meta_run.italic = True
                    meta_run.font.size = Pt(9)
                for line in str(item.get("bullets") or "").split("\n"):
                    line = line.strip(" •-\t")
                    if line:
                        doc.add_paragraph(line, style="List Bullet")
        elif key == "education":
            blocks = [item for item in data.get("education") or [] if entry_has_content(item)]
            if not blocks:
                continue
            _heading(doc, "Education", theme)
            for item in blocks:
                title = " • ".join(
                    part for part in (item.get("degree"), item.get("field")) if part
                ) or item.get("school")
                paragraph = doc.add_paragraph()
                run = paragraph.add_run(title or "")
                run.bold = True
                meta = "  |  ".join(
                    bit
                    for bit in (
                        item.get("school") if title != item.get("school") else "",
                        format_dates(item.get("start"), item.get("end")),
                        item.get("location") or "",
                        f"GPA {item['gpa']}" if item.get("gpa") else "",
                    )
                    if bit
                )
                if meta:
                    doc.add_paragraph(meta)
                if item.get("details"):
                    doc.add_paragraph(item["details"])
        elif key == "projects":
            blocks = [item for item in data.get("projects") or [] if entry_has_content(item)]
            if not blocks:
                continue
            _heading(doc, "Projects", theme)
            for item in blocks:
                paragraph = doc.add_paragraph()
                run = paragraph.add_run(item.get("name") or "Project")
                run.bold = True
                if item.get("tech"):
                    doc.add_paragraph(item["tech"])
                link_p = doc.add_paragraph()
                for label, url in (
                    ("GitHub", item.get("github")),
                    ("Live", item.get("live") or item.get("url")),
                ):
                    if url and str(url).startswith("http"):
                        _add_hyperlink(link_p, label + "  ", url)
                if item.get("description"):
                    doc.add_paragraph(item["description"])

    skills = skill_list(data)
    if skills:
        _heading(doc, "Skills", theme)
        doc.add_paragraph(", ".join(skills))

    certs = [item for item in data.get("certifications") or [] if entry_has_content(item)]
    if certs:
        _heading(doc, "Certifications", theme)
        for item in certs:
            line = " — ".join(
                part
                for part in (
                    item.get("name"),
                    " | ".join(part for part in (item.get("issuer"), item.get("year")) if part),
                )
                if part
            )
            if line:
                doc.add_paragraph(line, style="List Bullet")

    achievements = [
        item for item in data.get("achievements") or [] if entry_has_content(item)
    ]
    if achievements:
        _heading(doc, "Achievements", theme)
        for item in achievements:
            line = " — ".join(
                part for part in (item.get("title"), item.get("description")) if part
            )
            if line:
                doc.add_paragraph(line, style="List Bullet")

    languages = csv_list(data.get("languages"))
    if languages:
        _heading(doc, "Languages", theme)
        doc.add_paragraph(", ".join(languages))

    interests = csv_list(data.get("interests"))
    if interests:
        _heading(doc, "Interests", theme)
        doc.add_paragraph(", ".join(interests))

    for item in data.get("custom_sections") or []:
        if not entry_has_content(item):
            continue
        _heading(doc, item.get("title") or "Additional", theme)
        if item.get("content"):
            doc.add_paragraph(item["content"])

    doc.save(output_file)


def _heading(doc, text, theme):
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(11)
    accent = theme["accent"] if theme["accent"] != "#f8fafc" else "#0f172a"
    _set_run_color(run, accent)
    _add_bottom_border(paragraph, theme["rule"])


def _generate_legacy_docx(resume, output_file):
    doc = Document()
    doc.add_heading("Professional Resume", level=1)
    sections = resume["sections"]
    for section in sections.values():
        doc.add_heading(section["title"], level=2)
        content = section["content"]
        if isinstance(content, list):
            for item in content:
                if str(item).strip():
                    doc.add_paragraph(str(item), style="List Bullet")
        else:
            text = str(content).strip()
            if text and text != "Professional summary not available.":
                doc.add_paragraph(text)
    doc.save(output_file)
