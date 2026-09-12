from docx import Document


def create_docx(resume, output_path):

    doc = Document()

    doc.add_heading(resume["name"], level=1)

    doc.add_paragraph(resume["email"])
    doc.add_paragraph(resume["phone"])

    doc.add_heading("Professional Summary", level=2)
    doc.add_paragraph(resume["summary"])

    doc.add_heading("Skills", level=2)

    for skill in resume["skills"]:
        doc.add_paragraph(skill, style="List Bullet")

    doc.add_heading("Projects", level=2)

    for project in resume["projects"]:
        doc.add_paragraph(project, style="List Bullet")

    doc.add_heading("Education", level=2)

    for edu in resume["education"]:
        doc.add_paragraph(edu, style="List Bullet")

    doc.add_heading("Certifications", level=2)

    for cert in resume["certifications"]:
        doc.add_paragraph(cert, style="List Bullet")

    doc.save(output_path)